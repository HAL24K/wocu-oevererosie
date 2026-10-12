"""One slide per example location: aerial photo, scope regions coloured by VVR crossing year,
the VVR, measured bank lines and predicted bank lines (2027, 2031).

Run: ``PYTHONPATH=. uv run python scripts/slides/location_slides.py --run 20260907-near-fz
        --hybrid-gpkg …_nearest_20260710.gpkg --out docs/presentations/20261007_locaties.pptx``
"""

from __future__ import annotations

import argparse
import warnings
from pathlib import Path

import contextily as ctx
import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pyogrio
from matplotlib.colors import LinearSegmentedColormap, Normalize
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from pptx import Presentation
from pptx.util import Inches, Pt
from shapely.geometry import box

from src.pipeline.config import ExperimentConfig
from src.sources.geometry import LOCATION_ID, normalise_location_id

warnings.filterwarnings("ignore")
ap = argparse.ArgumentParser()
ap.add_argument("--run", required=True)
ap.add_argument("--hybrid-gpkg", required=True)
ap.add_argument("--out", required=True)
ap.add_argument("--half-width-m", type=float, default=260.0)
ap.add_argument("--experiments", help="CSV from scripts/evaluate/summarise_mask_runs.py; adds a results slide up front")
ap.add_argument("--takeaways", nargs="*", default=[], help="bullets under the results table")
args = ap.parse_args()

# (session, place, who, region, note)
LOCATIONS = [
    ("14 okt · Rijntakken", "Rechtsonder Wageningen", "Pam", "nederrijn_r_3200_3210", ""),
    ("14 okt · Rijntakken", "Onder Elst (Elster Buitenwaard)", "Pam", "nederrijn_r_4780_4790", ""),
    ("14 okt · Rijntakken", "Boven Zuilichem", "Pam", "rijn_l_8380_8390", ""),
    ("14 okt · Rijntakken", "Bij Loevestein", "Pam", "rijn_l_9150_9160", ""),
    ("14 okt · Rijntakken", "Boven Druten", "Pam", "rijn_l_4650_4660", ""),
    ("15 okt · Maas", "Broekhuizerweerd", "Rik", "maas3_l_8860_8870", ""),
    ("15 okt · Maas", "Zuidereiland (bij A77)", "Rik", "maas3_r_11660_11670", "10,6 m/jr voorspeld: controleren, mogelijk een SAM-artefact."),
    ("15 okt · Maas", "Balgooij", "Rik", "maas3_r_14450_14460", "Herstel dit jaar: ziet het model de ingreep?"),
    ("15 okt · Maas", "Rijkelse Bemden", "Rik", "maas3_r_5840_5850", "Herstel dit jaar: ziet het model de ingreep?"),
    ("15 okt · Maas", "Mussenwaard", "Rik", "maas3_r_18650_18660", "Herstel dit jaar: ziet het model de ingreep?"),
]
AERIAL = "https://service.pdok.nl/hwh/luchtfotorgb/wmts/v1_0/Actueel_orthoHR/EPSG:3857/{z}/{x}/{y}.jpeg"
WARN = LinearSegmentedColormap.from_list("warn", ["#d7191c", "#f46d43", "#fdae61", "#fee08b", "#ffff33"])
WARN_NORM = Normalize(2027, 2031)
FILL_ALPHA = 0.6

cfg = ExperimentConfig(experiment=args.run)
TRIAGE = cfg.exports_dir
scope = normalise_location_id(gpd.read_file(cfg.scope_gpkg, layer="vlakken_scope")).to_crs(28992)
vvr_pl = gpd.read_file(cfg.signalering_gpkg, layer="Vlak_vrije_ruimte_natuurvriendelijke_oever_pl").to_crs(28992)
vvr_ln = gpd.read_file(cfg.signalering_gpkg, layer="Vlak_vrije_ruimte_natuurvriendelijke_oever_ln").to_crs(28992)
cross = gpd.read_file(TRIAGE / f"vvr_kruising_{args.run}.gpkg", layer="vvr_kruising_vlakken").set_index(LOCATION_ID)
pred = gpd.read_file(TRIAGE / f"voorspelde_oeverlijnen_2026_2031_{args.run}.gpkg")
measured = normalise_location_id(pyogrio.read_dataframe(args.hybrid_gpkg, layer="lines"))
measured["year"] = pd.to_datetime(measured["date"]).dt.year
pref = normalise_location_id(pyogrio.read_dataframe(args.hybrid_gpkg, layer="model_preference", read_geometry=False))
pref = pref.set_index(LOCATION_ID)["model_preference"]
scope["kruisingsjaar"] = scope[LOCATION_ID].map(cross["kruisingsjaar"])
fig_dir = Path(args.out).with_suffix("")
fig_dir.mkdir(parents=True, exist_ok=True)


def frame(region: str):
    c = scope.loc[scope[LOCATION_ID] == region].geometry.iloc[0].centroid
    w = args.half_width_m
    return box(c.x - w * 1.25, c.y - w, c.x + w * 1.25, c.y + w)


def render(region: str, path: Path) -> dict:
    f = frame(region)
    fx0, fy0, fx1, fy1 = f.bounds
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.set_xlim(fx0, fx1)
    ax.set_ylim(fy0, fy1)
    ctx.add_basemap(ax, source=AERIAL, crs="EPSG:28992", attribution="Luchtfoto PDOK", zoom=17)
    sc = scope[scope.intersects(f)]
    warn = sc[sc["kruisingsjaar"].notna() & (sc["kruisingsjaar"] <= 2031)]
    for _, r in warn.iterrows():
        col = WARN(WARN_NORM(max(r["kruisingsjaar"], 2027)))
        gpd.GeoSeries([r.geometry]).plot(ax=ax, color=col, alpha=FILL_ALPHA, linewidth=0)
    sc.boundary.plot(ax=ax, color="white", linewidth=0.5, alpha=0.8)
    sc[sc[LOCATION_ID] == region].boundary.plot(ax=ax, color="white", linewidth=2.6)
    vp = vvr_pl[vvr_pl.intersects(f)]
    if len(vp):
        vp.boundary.plot(ax=ax, color="#00e676", linewidth=1.6, linestyle="--")
    vl = vvr_ln[vvr_ln.intersects(f)]
    if len(vl):
        vl.plot(ax=ax, color="#00e676", linewidth=2.4)
    m = measured[measured.intersects(f) & measured[LOCATION_ID].isin(sc[LOCATION_ID])]
    years = sorted(m["year"].unique())
    blues = plt.get_cmap("cool")
    for i, (y, g) in enumerate(m.groupby("year")):
        g.plot(ax=ax, color=blues(i / max(len(years) - 1, 1)), linewidth=1.3, alpha=0.9)
    p = pred[pred[LOCATION_ID].isin(sc[LOCATION_ID]) & pred["year"].isin([2027, 2031])]
    for y, g in p.groupby("year"):
        g.plot(ax=ax, color="black", linewidth=4.2, alpha=0.6)
        g.plot(ax=ax, color=WARN(WARN_NORM(y)), linewidth=2.4, linestyle=(0, (4, 2)))
    ax.set_axis_off()
    handles = [
        Patch(facecolor=WARN(WARN_NORM(2027)), alpha=FILL_ALPHA, label="raakt VVR ≤ 2027 (of al)"),
        Patch(facecolor=WARN(WARN_NORM(2029)), alpha=FILL_ALPHA, label="raakt VVR in 2029"),
        Patch(facecolor=WARN(WARN_NORM(2031)), alpha=FILL_ALPHA, label="raakt VVR in 2031"),
        Line2D([], [], color="#00e676", lw=2.4, label="signaleringslijn VVR"),
        Line2D([], [], color="#00e676", lw=1.6, ls="--", label="vlak vrije ruimte (VVR)"),
        Line2D([], [], color=blues(0.0), lw=1.5, label=f"gemeten oever {years[0] if years else ''}"),
        Line2D([], [], color=blues(1.0), lw=1.5, label=f"gemeten oever {years[-1] if years else ''}"),
        Line2D([], [], color=WARN(WARN_NORM(2027)), lw=2.4, ls="--", label="voorspelde oever 2027"),
        Line2D([], [], color=WARN(WARN_NORM(2031)), lw=2.4, ls="--", label="voorspelde oever 2031"),
        Line2D([], [], color="white", lw=2.6, label="deze scoperegio"),
    ]
    leg = ax.legend(handles=handles, loc="lower left", fontsize=8, framealpha=0.85)
    leg.get_frame().set_edgecolor("none")
    fig.savefig(path, dpi=160, bbox_inches="tight", pad_inches=0)
    plt.close(fig)
    pr = pred[(pred[LOCATION_ID] == region) & (pred["year"] == 2031)]
    mr = m[m[LOCATION_ID] == region]
    return {
        "v": float(pr["v_mean_m_per_yr"].iloc[0]) if len(pr) else np.nan,
        "cross": cross["kruisingsjaar"].get(region, np.nan),
        "pct": cross["pct_vvr_op_2026"].get(region, np.nan),
        "pref": pref.get(region, "?"),
        "years": sorted(mr["year"].unique()),
    }


def facts(place: str, who: str, region: str, note: str, d: dict) -> list[str]:
    source = {"height": "hoogtemodel", "segmentation": "SAM (satelliet)"}.get(d["pref"], d["pref"])
    lines = [
        f"Locatie: {place} (aangedragen door {who})",
        f"Scoperegio: {region}",
        f"Bron gekozen door hybride model: {source}",
        f"Gemeten jaren: {', '.join(map(str, d['years'])) or 'geen'}",
        f"Voorspelde verplaatsing: {d['v']:+.1f} m/jr (+ = landinwaarts)",
    ]
    if np.isnan(d["cross"]):
        lines.append("VVR: geen kruising voorspeld (of geen VVR-vlak)")
    elif d["cross"] <= 2026:
        lines.append(f"VVR: al bereikt · {d['pct']:.0f} % van de vrije ruimte gebruikt in 2026")
    else:
        lines.append(f"VVR: bereikt in {int(d['cross'])} · {d['pct']:.0f} % gebruikt in 2026")
    if note:
        lines.append(f"Let op: {note}")
    return lines


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
if args.experiments:
    ex = pd.read_csv(args.experiments)
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.8)).text_frame
    tb.text = "Hybride model met onze eigen instellingen: maskers op de SAM-lijnen"
    tb.paragraphs[0].runs[0].font.size = Pt(24)
    tb.paragraphs[0].runs[0].font.bold = True
    sub = tb.add_paragraph()
    sub.text = "Masker toegepast na Lukes nearest-filter, vóór de puntentelling · zelfde hoogtepunten, pipeline en bevroren testset · fout op 5 m-strookjes (m/jr)"
    sub.runs[0].font.size = Pt(13)
    cols = [
        ("variant", "variant", "{}"),
        ("SAM→height", "SAM → hoogte", "{:,.0f}"),
        ("height→SAM", "hoogte → SAM", "{:,.0f}"),
        ("MAE all", "MAE alle", "{:.2f}"),
        ("skill all", "skill alle", "{:+.0%}"),
        ("MAE VVR", "MAE VVR", "{:.2f}"),
        ("skill VVR", "skill VVR", "{:+.0%}"),
        ("MAE height-preferred", "MAE hoogte-regio's", "{:.2f}"),
        ("MAE SAM-preferred", "MAE SAM-regio's", "{:.2f}"),
        ("VVR crossing by 2027", "VVR geraakt ≤ 2027", "{:,.0f}"),
    ]
    tbl = s.shapes.add_table(len(ex) + 1, len(cols), Inches(0.4), Inches(1.4), Inches(12.5), Inches(0.4) * (len(ex) + 1)).table
    for j, (_, head, _) in enumerate(cols):
        tbl.cell(0, j).text = head
    for i, r in ex.iterrows():
        for j, (key, _, fmt) in enumerate(cols):
            tbl.cell(i + 1, j).text = fmt.format(r[key])
    for row in tbl.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
    body = s.shapes.add_textbox(Inches(0.4), Inches(1.6) + Inches(0.4) * (len(ex) + 1), Inches(12.5), Inches(3)).text_frame
    body.word_wrap = True
    for i, line in enumerate(args.takeaways):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = f"• {line}"
        para.runs[0].font.size = Pt(15)
        para.space_after = Pt(8)

for session, place, who, region, note in LOCATIONS:
    png = fig_dir / f"{region}.png"
    d = render(region, png)
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.8)).text_frame
    tb.text = f"{place}"
    tb.paragraphs[0].runs[0].font.size = Pt(26)
    tb.paragraphs[0].runs[0].font.bold = True
    p2 = tb.add_paragraph()
    p2.text = session
    p2.runs[0].font.size = Pt(14)
    s.shapes.add_picture(str(png), Inches(0.4), Inches(1.2), height=Inches(6.0))
    body = s.shapes.add_textbox(Inches(8.2), Inches(1.4), Inches(4.8), Inches(5.6)).text_frame
    body.word_wrap = True
    for i, line in enumerate(facts(place, who, region, note, d)):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = line
        para.runs[0].font.size = Pt(13)
        para.space_after = Pt(8)
    print(f"{place}: {region} · v {d['v']:+.2f} · kruising {d['cross']} · {d['pref']}", flush=True)
prs.save(args.out)
print("→", args.out)
